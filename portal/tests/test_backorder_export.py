from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from openpyxl import load_workbook

from backorders.models import BackorderLine, BackorderOrderFolder
from hospitals.models import Hospital
from orders.models import Order
from portal.services.backorder_export_service import build_backorder_xlsx
from portal.services.backorder_portal_service import build_backorder_queryset
from shipments.models import ShipmentBatch


XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
EXPECTED_DETAIL_HEADERS = [
    "医院订单",
    "下一批编号",
    "医院",
    "产品号",
    "订单数量",
    "已发数量",
    "待补发",
]
EXPECTED_SUMMARY_HEADERS = [
    "产品号",
    "总待补发",
    "涉及订单数",
    "订单分配",
]


class BackorderExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="backorder-export-staff",
            password="test-password",
            is_staff=True,
        )
        self.hospital = Hospital.objects.create(name="Canonical Hospital")
        self.primary_line = self.create_line(
            order_number="00146769",
            hospital_name="Fallback Hospital",
            hospital=self.hospital,
            product_code="00042-A",
            requested_quantity=7,
            shipped_quantity=2,
            remaining_quantity=5,
        )
        self.fallback_line = self.create_line(
            order_number="FALLBACK-ORDER",
            hospital_name="Fallback Only Hospital",
            product_code="FALLBACK-PRODUCT",
            requested_quantity=9,
            shipped_quantity=4,
            remaining_quantity=5,
        )
        self.completed_line = self.create_line(
            order_number="COMPLETED-ORDER",
            hospital_name="Completed Hospital",
            product_code="COMPLETED-PRODUCT",
            requested_quantity=4,
            shipped_quantity=4,
            remaining_quantity=0,
        )
        self.url = reverse("portal:backorder_export_xlsx")

    def create_line(
        self,
        *,
        order_number=None,
        hospital_name="Test Hospital",
        product_code,
        requested_quantity,
        shipped_quantity,
        remaining_quantity,
        hospital=None,
        order=None,
    ):
        if order is None:
            order = Order.objects.create(
                bon_de_commande=order_number,
                hospital_name=hospital_name,
                hospital=hospital,
                hospital_order_pdf=f"hospital_orders/{order_number}.pdf",
                created_by=self.user,
            )
            folder = BackorderOrderFolder.objects.create(order=order)
        else:
            folder = order.backorder_folder
        return BackorderLine.objects.create(
            order_folder=folder,
            order=order,
            product_code=product_code,
            requested_quantity=requested_quantity,
            shipped_quantity=shipped_quantity,
            remaining_quantity=remaining_quantity,
        )

    def create_batch(self, order, batch_number):
        return ShipmentBatch.objects.create(
            order=order,
            batch_number=batch_number,
            month_key="2026-09",
        )

    def export_workbook(self, query=""):
        self.client.force_login(self.user)
        response = self.client.get(f"{self.url}{query}")
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        return response, workbook, workbook["待补发明细"]

    def test_export_requires_authenticated_staff_user(self):
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 302)

    def test_workbook_has_both_sheets_and_expected_layout(self):
        response, workbook, detail_sheet = self.export_workbook("?q=00146769")
        summary_sheet = workbook["产品补货汇总"]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertIn(".xlsx", response["Content-Disposition"])
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])
        self.assertEqual(
            workbook.sheetnames,
            ["待补发明细", "产品补货汇总"],
        )
        self.assertEqual(
            [cell.value for cell in detail_sheet[1]],
            EXPECTED_DETAIL_HEADERS,
        )
        self.assertEqual(
            [cell.value for cell in summary_sheet[1]],
            EXPECTED_SUMMARY_HEADERS,
        )
        self.assertEqual(detail_sheet.freeze_panes, "A2")
        self.assertEqual(summary_sheet.freeze_panes, "A2")
        self.assertEqual(detail_sheet.auto_filter.ref, "A1:G2")
        self.assertEqual(summary_sheet.auto_filter.ref, "A1:D2")
        self.assertTrue(all(cell.font.bold for cell in detail_sheet[1]))
        self.assertTrue(
            all(
                cell.alignment.horizontal == "center"
                for cell in detail_sheet[1]
            )
        )
        self.assertEqual(
            [
                detail_sheet.column_dimensions[column].width
                for column in "ABCDEFG"
            ],
            [18, 20, 35, 20, 12, 12, 12],
        )

    def test_same_order_shows_order_fields_only_on_first_product_row(self):
        order = self.primary_line.order
        self.create_line(
            order=order,
            product_code="00042-B",
            requested_quantity=3,
            shipped_quantity=1,
            remaining_quantity=2,
        )
        self.create_line(
            order=order,
            product_code="00042-C",
            requested_quantity=6,
            shipped_quantity=0,
            remaining_quantity=6,
        )

        _response, _workbook, sheet = self.export_workbook("?q=00146769")

        self.assertEqual(
            [cell.value for cell in sheet[2][:4]],
            ["00146769", "00146769", "Canonical Hospital", "00042-A"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[3][:4]],
            [None, None, None, "00042-B"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[4][:4]],
            [None, None, None, "00042-C"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[5][3:7]],
            ["订单小计", 16, 3, 13],
        )
        self.assertTrue(all(cell.font.bold for cell in sheet[5]))
        self.assertEqual(
            [cell.value for cell in sheet[6][3:7]],
            ["总计", 16, 3, 13],
        )

    def test_next_reference_uses_max_batch_number_plus_one(self):
        order = self.primary_line.order
        self.create_batch(order, 1)
        self.create_batch(order, 2)

        _response, _workbook, sheet = self.export_workbook("?q=00146769")

        self.assertEqual(sheet["B2"].value, "00146769-B3")

    def test_next_reference_uses_true_max_not_batch_count(self):
        order = self.primary_line.order
        self.create_batch(order, 1)
        self.create_batch(order, 8)

        _response, _workbook, sheet = self.export_workbook("?q=00146769")

        self.assertEqual(sheet["B2"].value, "00146769-B9")

    def test_next_reference_without_batches_is_original_bon(self):
        _response, _workbook, sheet = self.export_workbook("?q=00146769")

        self.assertEqual(sheet["B2"].value, "00146769")
        self.assertNotEqual(sheet["B2"].value, "00146769-B1")

    def test_export_maps_fields_and_preserves_text_and_numeric_types(self):
        _response, _workbook, sheet = self.export_workbook("?q=00146769")

        self.assertEqual(
            [cell.value for cell in sheet[2]],
            [
                "00146769",
                "00146769",
                "Canonical Hospital",
                "00042-A",
                7,
                2,
                5,
            ],
        )
        self.assertEqual(
            [cell.data_type for cell in sheet[2][:4]],
            ["s", "s", "s", "s"],
        )
        self.assertEqual(
            [cell.data_type for cell in sheet[2][4:7]],
            ["n", "n", "n"],
        )
        self.assertEqual(
            [cell.number_format for cell in sheet[2][:4]],
            ["@", "@", "@", "@"],
        )
        self.assertEqual(
            [cell.number_format for cell in sheet[2][4:7]],
            ["0", "0", "0"],
        )

    def test_export_falls_back_to_order_hospital_name(self):
        _response, _workbook, sheet = self.export_workbook(
            "?q=FALLBACK-ORDER"
        )

        self.assertEqual(sheet["C2"].value, "Fallback Only Hospital")

    def test_detail_order_and_global_totals_are_correct(self):
        _response, _workbook, sheet = self.export_workbook()

        subtotal_rows = [
            row
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[3] == "订单小计"
        ]
        self.assertEqual(
            [(row[4], row[5], row[6]) for row in subtotal_rows],
            [(7, 2, 5), (9, 4, 5)],
        )
        self.assertEqual(
            [cell.value for cell in sheet[sheet.max_row][3:7]],
            ["总计", 16, 6, 10],
        )

    def test_product_summary_aggregates_orders_and_allocation_text(self):
        second_line = self.create_line(
            order_number="00150000",
            product_code="00042-A",
            requested_quantity=4,
            shipped_quantity=1,
            remaining_quantity=3,
        )
        self.create_batch(self.primary_line.order, 2)
        self.create_batch(second_line.order, 1)

        _response, workbook, _detail = self.export_workbook(
            "?q=00042-A&status=active"
        )
        summary = workbook["产品补货汇总"]

        self.assertEqual(
            [cell.value for cell in summary[2]],
            [
                "00042-A",
                8,
                2,
                "00146769-B3 ×5; 00150000-B2 ×3",
            ],
        )

    def test_detail_and_product_summary_remaining_totals_reconcile(self):
        _response, workbook, detail = self.export_workbook()
        summary = workbook["产品补货汇总"]

        detail_total = next(
            row[6]
            for row in detail.iter_rows(min_row=2, values_only=True)
            if row[3] == "总计"
        )
        summary_products_total = sum(
            row[1]
            for row in summary.iter_rows(min_row=2, values_only=True)
            if row[0] not in {"产品总类数", "产品总计"}
        )
        summary_total = next(
            row[1]
            for row in summary.iter_rows(min_row=2, values_only=True)
            if row[0] == "产品总计"
        )

        self.assertEqual(detail_total, summary_products_total)
        self.assertEqual(detail_total, summary_total)

    def test_filtered_queryset_limits_both_sheets_and_totals(self):
        _response, workbook, detail = self.export_workbook(
            "?q=00042-A&status=active"
        )
        summary = workbook["产品补货汇总"]

        self.assertEqual(detail["D2"].value, "00042-A")
        self.assertEqual(detail["D4"].value, "总计")
        self.assertEqual(detail["G4"].value, 5)
        self.assertEqual(summary["A2"].value, "00042-A")
        self.assertEqual(summary["B2"].value, 5)
        self.assertEqual(summary["A4"].value, "产品总计")
        self.assertEqual(summary["B4"].value, 5)

    def test_default_export_only_includes_active_remaining_lines(self):
        _response, _workbook, sheet = self.export_workbook()

        exported_orders = {
            row[0]
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[0]
        }
        self.assertEqual(exported_orders, {"00146769", "FALLBACK-ORDER"})
        self.assertNotIn(
            self.completed_line.order.bon_de_commande,
            exported_orders,
        )

    def test_export_ignores_page_and_exports_all_filtered_results(self):
        bulk_order = Order.objects.create(
            bon_de_commande="BULK-ORDER",
            hospital_name="Bulk Hospital",
            hospital_order_pdf="hospital_orders/bulk.pdf",
            created_by=self.user,
        )
        folder = BackorderOrderFolder.objects.create(order=bulk_order)
        BackorderLine.objects.bulk_create(
            [
                BackorderLine(
                    order_folder=folder,
                    order=bulk_order,
                    product_code=f"BULK-{index:03d}",
                    requested_quantity=2,
                    shipped_quantity=1,
                    remaining_quantity=1,
                )
                for index in range(301)
            ]
        )

        _response, workbook, sheet = self.export_workbook(
            "?q=BULK-ORDER&page=99"
        )

        self.assertEqual(sheet.max_row, 304)
        self.assertEqual(workbook["产品补货汇总"].max_row, 304)

    def test_empty_export_has_headers_and_zero_totals(self):
        _response, workbook, detail = self.export_workbook(
            "?q=NO-SUCH-BACKORDER"
        )
        summary = workbook["产品补货汇总"]

        self.assertEqual(
            [cell.value for cell in detail[1]],
            EXPECTED_DETAIL_HEADERS,
        )
        self.assertEqual(
            [cell.value for cell in detail[2][3:7]],
            ["总计", 0, 0, 0],
        )
        self.assertEqual(detail.auto_filter.ref, "A1:G1")
        self.assertEqual(summary["A2"].value, "产品总类数")
        self.assertEqual(summary["B2"].value, 0)
        self.assertEqual(summary["A3"].value, "产品总计")
        self.assertEqual(summary["B3"].value, 0)
        self.assertEqual(summary.auto_filter.ref, "A1:D1")

    def test_formula_like_text_is_stored_as_text(self):
        dangerous_hospital = Hospital.objects.create(name="-Dangerous Hospital")
        self.create_line(
            order_number="=1+1",
            hospital_name="Unused",
            hospital=dangerous_hospital,
            product_code="@SUM(A1:A2)",
            requested_quantity=1,
            shipped_quantity=0,
            remaining_quantity=1,
        )

        _response, _workbook, sheet = self.export_workbook("?q=%3D1%2B1")

        self.assertEqual(
            [sheet.cell(row=2, column=column).value for column in range(1, 5)],
            ["=1+1", "=1+1", "-Dangerous Hospital", "@SUM(A1:A2)"],
        )
        self.assertEqual(
            [
                sheet.cell(row=2, column=column).data_type
                for column in range(1, 5)
            ],
            ["s", "s", "s", "s"],
        )

    def test_export_uses_constant_query_count_for_batch_lookup(self):
        queryset = build_backorder_queryset({"status": "active"})

        with self.assertNumQueries(2):
            build_backorder_xlsx(queryset)

    def test_list_export_link_preserves_filters_but_drops_page(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("portal:backorder_list"),
            {
                "q": "00146769",
                "status": "active",
                "hospital": "Canonical",
                "page": "4",
            },
        )

        self.assertEqual(
            response.context["export_query_string"],
            "q=00146769&status=active&hospital=Canonical",
        )
