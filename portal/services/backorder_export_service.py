from collections import defaultdict
from io import BytesIO

from django.db.models import Max
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from shipments.models import ShipmentBatch


BACKORDER_EXPORT_HEADERS = [
    "医院订单",
    "下一批编号",
    "医院",
    "产品号",
    "订单数量",
    "已发数量",
    "待补发",
]

BACKORDER_EXPORT_COLUMN_WIDTHS = {
    "A": 18,
    "B": 20,
    "C": 35,
    "D": 20,
    "E": 12,
    "F": 12,
    "G": 12,
}

PRODUCT_SUMMARY_HEADERS = [
    "产品号",
    "总待补发",
    "涉及订单数",
    "订单分配",
]

PRODUCT_SUMMARY_COLUMN_WIDTHS = {
    "A": 20,
    "B": 14,
    "C": 14,
    "D": 60,
}

SUMMARY_FILL = PatternFill(fill_type="solid", fgColor="E7EEF8")
SUMMARY_BORDER = Border(top=Side(style="thin", color="A6A6A6"))


def _set_text_cell(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.number_format = "@"


def _set_quantity_cell(cell, value):
    cell.value = int(value)
    cell.number_format = "0"


def _style_header(worksheet):
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def _style_summary_row(worksheet, row_number):
    for cell in worksheet[row_number]:
        cell.font = Font(bold=True)
        cell.fill = SUMMARY_FILL
        cell.border = SUMMARY_BORDER
        cell.alignment = Alignment(vertical="center")


def _next_references(lines):
    orders = {
        line.order_id: line.order.bon_de_commande
        for line in lines
    }
    max_batches = {
        row["order_id"]: row["max_batch_number"]
        for row in ShipmentBatch.objects.filter(
            order_id__in=orders,
        )
        .values("order_id")
        .annotate(max_batch_number=Max("batch_number"))
    }

    return {
        order_id: (
            f"{bon_de_commande}-B{max_batches[order_id] + 1}"
            if order_id in max_batches
            else bon_de_commande
        )
        for order_id, bon_de_commande in orders.items()
    }


def _hospital_name(order):
    if order.hospital:
        return order.hospital.name
    return order.hospital_name or ""


def _build_detail_sheet(workbook, lines, next_references):
    worksheet = workbook.active
    worksheet.title = "待补发明细"
    worksheet.append(BACKORDER_EXPORT_HEADERS)
    worksheet.freeze_panes = "A2"
    _style_header(worksheet)

    grand_totals = [0, 0, 0]
    current_order_id = None
    order_totals = [0, 0, 0]
    last_detail_row = 1

    def append_subtotal():
        if current_order_id is None:
            return
        row_number = worksheet.max_row + 1
        _set_text_cell(worksheet.cell(row=row_number, column=4), "订单小计")
        for column, quantity in enumerate(order_totals, start=5):
            _set_quantity_cell(
                worksheet.cell(row=row_number, column=column),
                quantity,
            )
        _style_summary_row(worksheet, row_number)

    for line in lines:
        if line.order_id != current_order_id:
            append_subtotal()
            current_order_id = line.order_id
            order_totals = [0, 0, 0]
            first_order_row = True
        else:
            first_order_row = False

        row_number = worksheet.max_row + 1
        if first_order_row:
            _set_text_cell(
                worksheet.cell(row=row_number, column=1),
                line.order.bon_de_commande,
            )
            _set_text_cell(
                worksheet.cell(row=row_number, column=2),
                next_references[line.order_id],
            )
            _set_text_cell(
                worksheet.cell(row=row_number, column=3),
                _hospital_name(line.order),
            )
        else:
            for column in range(1, 4):
                _set_text_cell(worksheet.cell(row=row_number, column=column), "")

        _set_text_cell(
            worksheet.cell(row=row_number, column=4),
            line.product_code,
        )
        quantities = (
            line.requested_quantity,
            line.shipped_quantity,
            line.remaining_quantity,
        )
        for index, quantity in enumerate(quantities):
            _set_quantity_cell(
                worksheet.cell(row=row_number, column=index + 5),
                quantity,
            )
            order_totals[index] += int(quantity)
            grand_totals[index] += int(quantity)

        for cell in worksheet[row_number]:
            cell.alignment = Alignment(vertical="center")
        last_detail_row = row_number

    append_subtotal()

    total_row = worksheet.max_row + 1
    _set_text_cell(worksheet.cell(row=total_row, column=4), "总计")
    for column, quantity in enumerate(grand_totals, start=5):
        _set_quantity_cell(
            worksheet.cell(row=total_row, column=column),
            quantity,
        )
    _style_summary_row(worksheet, total_row)

    worksheet.auto_filter.ref = f"A1:G{last_detail_row}"
    for column, width in BACKORDER_EXPORT_COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    return grand_totals


def _build_product_summary_sheet(
    workbook,
    lines,
    next_references,
    detail_remaining_total,
):
    worksheet = workbook.create_sheet("产品补货汇总")
    worksheet.append(PRODUCT_SUMMARY_HEADERS)
    worksheet.freeze_panes = "A2"
    _style_header(worksheet)

    product_data = defaultdict(
        lambda: {"remaining": 0, "orders": defaultdict(int)}
    )
    order_sort_keys = {}

    for line in lines:
        quantity = int(line.remaining_quantity)
        product_data[line.product_code]["remaining"] += quantity
        product_data[line.product_code]["orders"][line.order_id] += quantity
        order_sort_keys[line.order_id] = (
            str(line.order.bon_de_commande),
            next_references[line.order_id],
            line.order_id,
        )

    for product_code in sorted(product_data):
        data = product_data[product_code]
        row_number = worksheet.max_row + 1
        _set_text_cell(
            worksheet.cell(row=row_number, column=1),
            product_code,
        )
        _set_quantity_cell(
            worksheet.cell(row=row_number, column=2),
            data["remaining"],
        )
        _set_quantity_cell(
            worksheet.cell(row=row_number, column=3),
            len(data["orders"]),
        )
        allocation = "; ".join(
            f"{next_references[order_id]} ×{data['orders'][order_id]}"
            for order_id in sorted(data["orders"], key=order_sort_keys.get)
        )
        _set_text_cell(
            worksheet.cell(row=row_number, column=4),
            allocation,
        )
        for cell in worksheet[row_number]:
            cell.alignment = Alignment(vertical="center")

    last_product_row = worksheet.max_row
    product_count_row = worksheet.max_row + 1
    _set_text_cell(
        worksheet.cell(row=product_count_row, column=1),
        "产品总类数",
    )
    _set_quantity_cell(
        worksheet.cell(row=product_count_row, column=2),
        len(product_data),
    )
    _style_summary_row(worksheet, product_count_row)

    total_row = worksheet.max_row + 1
    _set_text_cell(worksheet.cell(row=total_row, column=1), "产品总计")
    _set_quantity_cell(
        worksheet.cell(row=total_row, column=2),
        detail_remaining_total,
    )
    _style_summary_row(worksheet, total_row)

    worksheet.auto_filter.ref = f"A1:D{last_product_row}"
    for column, width in PRODUCT_SUMMARY_COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width


def build_backorder_xlsx(queryset):
    lines = list(
        queryset.order_by(
            "order__bon_de_commande",
            "product_code",
            "id",
        )
    )
    next_references = _next_references(lines)

    workbook = Workbook()
    grand_totals = _build_detail_sheet(
        workbook,
        lines,
        next_references,
    )
    _build_product_summary_sheet(
        workbook,
        lines,
        next_references,
        detail_remaining_total=grand_totals[2],
    )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
