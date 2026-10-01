from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import (
    get_user_model,
)
from django.test import TestCase

from openpyxl import load_workbook

from documents.models import (
    GeneratedDocument,
)
from factories.models import Factory
from orders.models import Order
from settlements.models import (
    PaymentTransaction,
    SettlementAccount,
)


class SettlementFinanceExportTests(
    TestCase
):
    def setUp(self):
        self.user = (
            get_user_model()
            .objects.create_user(
                username=(
                    "finance-export-test"
                ),
                password="test-password",
                is_staff=True,
            )
        )

        self.factory = (
            Factory.objects.create(
                name="EXPORT TEST FACTORY",
                short_name="ETF",
            )
        )

        self.alpha_order = (
            self.create_financial_order(
                order_number=(
                    "EXPORT-ORDER-ALPHA"
                ),
                hospital_name=(
                    "EXPORT ALPHA HOSPITAL"
                ),
                invoice_number=(
                    "EXPORT-INV-ALPHA"
                ),
                po_number=(
                    "EXPORT-PO-ALPHA"
                ),
                sales=Decimal(
                    "1000.00"
                ),
                purchases=Decimal(
                    "400.00"
                ),
            )
        )

        self.beta_order = (
            self.create_financial_order(
                order_number=(
                    "EXPORT-ORDER-BETA"
                ),
                hospital_name=(
                    "EXPORT BETA HOSPITAL"
                ),
                invoice_number=(
                    "EXPORT-INV-BETA"
                ),
                po_number=(
                    "EXPORT-PO-BETA"
                ),
                sales=Decimal(
                    "2000.00"
                ),
                purchases=Decimal(
                    "800.00"
                ),
            )
        )

    def create_financial_order(
        self,
        *,
        order_number,
        hospital_name,
        invoice_number,
        po_number,
        sales,
        purchases,
    ):
        order = Order.objects.create(
            bon_de_commande=order_number,
            hospital_name=hospital_name,
            hospital_order_pdf=(
                "hospital_orders/"
                f"{order_number}.pdf"
            ),
            factory=self.factory,
            created_by=self.user,
        )

        invoice = (
            GeneratedDocument
            .objects.create(
                order=order,
                document_type=(
                    GeneratedDocument
                    .DocumentType
                    .HOSPITAL_INVOICE
                ),
                document_number=(
                    invoice_number
                ),
                generated_by=self.user,
                source_data={},
            )
        )

        po = (
            GeneratedDocument
            .objects.create(
                order=order,
                document_type=(
                    GeneratedDocument
                    .DocumentType
                    .FACTORY_PO
                ),
                document_number=po_number,
                generated_by=self.user,
                source_data={},
            )
        )

        receivable = (
            SettlementAccount
            .objects.create(
                document=invoice,
                direction=(
                    SettlementAccount
                    .Direction
                    .RECEIVABLE
                ),
                counterparty_name=(
                    hospital_name
                ),
                issue_date=date(
                    2026,
                    7,
                    1,
                ),
                due_date=date(
                    2026,
                    7,
                    31,
                ),
                currency="EUR",
                original_amount=sales,
            )
        )

        payable = (
            SettlementAccount
            .objects.create(
                document=po,
                direction=(
                    SettlementAccount
                    .Direction
                    .PAYABLE
                ),
                counterparty_name=(
                    self.factory.name
                ),
                issue_date=date(
                    2026,
                    7,
                    1,
                ),
                due_date=date(
                    2026,
                    8,
                    31,
                ),
                currency="EUR",
                original_amount=(
                    purchases
                ),
            )
        )

        PaymentTransaction.objects.create(
            account=receivable,
            payment_date=date(
                2026,
                7,
                10,
            ),
            amount=(
                sales / Decimal("4")
            ),
            method=(
                PaymentTransaction
                .Method
                .BANK_TRANSFER
            ),
            reference=(
                f"RECEIPT-{order_number}"
            ),
            created_by=self.user,
        )

        PaymentTransaction.objects.create(
            account=payable,
            payment_date=date(
                2026,
                7,
                11,
            ),
            amount=(
                purchases / Decimal("4")
            ),
            method=(
                PaymentTransaction
                .Method
                .BANK_TRANSFER
            ),
            reference=(
                f"PAYMENT-{order_number}"
            ),
            created_by=self.user,
        )

        return order

    def workbook_values(
        self,
        worksheet,
    ):
        return [
            value
            for row in worksheet.iter_rows(
                values_only=True
            )
            for value in row
            if value is not None
        ]

    def test_login_is_required(self):
        response = self.client.get(
            "/portal/finance/export.xlsx"
        )

        self.assertEqual(
            response.status_code,
            302,
        )

    def test_export_returns_xlsx(self):
        self.client.force_login(
            self.user
        )

        response = self.client.get(
            "/portal/finance/export.xlsx"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response["Content-Type"],
            (
                "application/vnd."
                "openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        self.assertTrue(
            response.content.startswith(
                b"PK"
            )
        )

        workbook = load_workbook(
            BytesIO(response.content),
            data_only=False,
        )

        self.assertEqual(
            workbook.sheetnames,
            [
                "摘要",
                "订单明细",
                "应收应付",
                "收付款流水",
            ],
        )

    def test_export_contains_finance_data(
        self,
    ):
        self.client.force_login(
            self.user
        )

        response = self.client.get(
            "/portal/finance/export.xlsx"
        )

        workbook = load_workbook(
            BytesIO(response.content),
            data_only=False,
        )

        order_values = (
            self.workbook_values(
                workbook["订单明细"]
            )
        )

        account_values = (
            self.workbook_values(
                workbook["应收应付"]
            )
        )

        transaction_values = (
            self.workbook_values(
                workbook["收付款流水"]
            )
        )

        self.assertIn(
            "EXPORT-ORDER-ALPHA",
            order_values,
        )

        self.assertIn(
            "EXPORT-INV-ALPHA",
            account_values,
        )

        self.assertIn(
            (
                "RECEIPT-"
                "EXPORT-ORDER-ALPHA"
            ),
            transaction_values,
        )

    def test_export_uses_order_filter(
        self,
    ):
        self.client.force_login(
            self.user
        )

        response = self.client.get(
            "/portal/finance/export.xlsx",
            {
                "order": "ALPHA",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        workbook = load_workbook(
            BytesIO(response.content),
            data_only=False,
        )

        values = self.workbook_values(
            workbook["订单明细"]
        )

        self.assertIn(
            "EXPORT-ORDER-ALPHA",
            values,
        )

        self.assertNotIn(
            "EXPORT-ORDER-BETA",
            values,
        )

    def test_invalid_date_returns_400(
        self,
    ):
        self.client.force_login(
            self.user
        )

        response = self.client.get(
            "/portal/finance/export.xlsx",
            {
                "date_from": (
                    "not-a-date"
                ),
            },
        )

        self.assertEqual(
            response.status_code,
            400,
        )
