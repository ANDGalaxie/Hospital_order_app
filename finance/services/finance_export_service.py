from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from finance.services.settlement_finance_service import (
    ZERO,
    build_filtered_account_queryset,
    build_settlement_finance_dashboard_data,
    money,
)
from settlements.models import (
    PaymentTransaction,
    SettlementAccount,
)


TITLE_FILL = PatternFill(
    "solid",
    fgColor="172554",
)

SECTION_FILL = PatternFill(
    "solid",
    fgColor="1E3A5F",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="DCE6F1",
)

INPUT_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
)

HEADER_FONT = Font(
    color="0F172A",
    bold=True,
)

INPUT_FONT = Font(
    color="0000FF",
)

FORMULA_FONT = Font(
    color="000000",
)

THIN_GREY = Side(
    style="thin",
    color="D9E2F3",
)

TABLE_BORDER = Border(
    bottom=THIN_GREY,
)

MONEY_FORMAT = (
    '€ #,##0.00;'
    '[Red](€ #,##0.00);'
    '-'
)

PERCENT_FORMAT = (
    '0.00%;'
    '[Red](0.00%);'
    '-'
)

DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"


def excel_datetime(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = (
                timezone.localtime(value)
                .replace(tzinfo=None)
            )

    return value


def user_display(user):
    if user is None:
        return ""

    full_name = str(
        user.get_full_name() or ""
    ).strip()

    return (
        full_name
        or user.get_username()
    )


def set_sheet_title(
    worksheet,
    title,
    *,
    end_column,
):
    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    cell = worksheet.cell(
        row=1,
        column=1,
        value=title,
    )

    cell.fill = TITLE_FILL
    cell.font = Font(
        color="FFFFFF",
        bold=True,
        size=16,
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 28


def style_header_row(
    worksheet,
    row_number,
    column_count,
):
    for column in range(
        1,
        column_count + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column,
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            top=THIN_GREY,
            bottom=THIN_GREY,
        )

    worksheet.row_dimensions[
        row_number
    ].height = 28


def apply_table_borders(
    worksheet,
    *,
    start_row,
    end_row,
    column_count,
):
    if end_row < start_row:
        return

    for row in worksheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=1,
        max_col=column_count,
    ):
        for cell in row:
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def set_column_widths(
    worksheet,
    widths,
):
    for column_index, width in enumerate(
        widths,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width


def write_summary_sheet(
    workbook,
    *,
    dashboard_data,
    filters,
):
    worksheet = workbook.active
    worksheet.title = "摘要"
    worksheet.sheet_view.showGridLines = False

    set_sheet_title(
        worksheet,
        "Acoeur 财务数据导出",
        end_column=4,
    )

    generated_at = excel_datetime(
        dashboard_data["generated_at"]
    )

    metadata = [
        (
            "导出时间",
            generated_at,
        ),
        (
            "记账币种",
            dashboard_data[
                "reporting_currency"
            ],
        ),
        (
            "开立日期从",
            filters.get("date_from"),
        ),
        (
            "开立日期至",
            filters.get("date_to"),
        ),
        (
            "医院筛选",
            filters.get(
                "hospital_query"
            )
            or "全部",
        ),
        (
            "工厂筛选",
            filters.get(
                "factory_query"
            )
            or "全部",
        ),
        (
            "订单号筛选",
            filters.get(
                "order_query"
            )
            or "全部",
        ),
        (
            "结算状态筛选",
            filters.get("status")
            or "全部",
        ),
    ]

    worksheet[
        "A3"
    ] = "导出条件"

    worksheet["A3"].fill = SECTION_FILL
    worksheet["A3"].font = WHITE_FONT

    worksheet.merge_cells(
        "A3:D3"
    )

    row_number = 4

    for label, value in metadata:
        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = HEADER_FONT

        value_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        value_cell.fill = INPUT_FILL
        value_cell.font = INPUT_FONT

        if isinstance(
            value,
            datetime,
        ):
            value_cell.number_format = (
                DATETIME_FORMAT
            )

        elif isinstance(
            value,
            date,
        ):
            value_cell.number_format = (
                DATE_FORMAT
            )

        worksheet.merge_cells(
            start_row=row_number,
            start_column=2,
            end_row=row_number,
            end_column=4,
        )

        row_number += 1

    row_number += 1

    worksheet.cell(
        row=row_number,
        column=1,
        value="核心指标",
    )

    worksheet.cell(
        row=row_number,
        column=1,
    ).fill = SECTION_FILL

    worksheet.cell(
        row=row_number,
        column=1,
    ).font = WHITE_FONT

    worksheet.merge_cells(
        start_row=row_number,
        start_column=1,
        end_row=row_number,
        end_column=4,
    )

    row_number += 1

    headers = [
        "指标",
        "金额 / 比例",
        "账户数量",
        "流水数量",
    ]

    for index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=row_number,
            column=index,
            value=header,
        )

    style_header_row(
        worksheet,
        row_number,
        len(headers),
    )

    summary = dashboard_data["summary"]

    kpi_rows = [
        (
            "医院开票销售额",
            summary["sales_total"],
        ),
        (
            "工厂采购额",
            summary["purchase_total"],
        ),
        (
            "预计毛利润",
            summary["gross_profit"],
        ),
        (
            "预计毛利率",
            summary["gross_margin"],
        ),
        (
            "医院已收金额",
            summary["receipt_total"],
        ),
        (
            "工厂已付金额",
            summary["payment_total"],
        ),
        (
            "实际现金净流入",
            summary["cash_net_inflow"],
        ),
        (
            "医院待收余额",
            summary[
                "receivable_remaining"
            ],
        ),
        (
            "工厂待付余额",
            summary[
                "payable_remaining"
            ],
        ),
        (
            "逾期应收金额",
            summary[
                "overdue_receivable_total"
            ],
        ),
        (
            "30 天内到期应收",
            summary[
                "due_soon_receivable_total"
            ],
        ),
    ]

    first_kpi_row = row_number + 1

    for label, value in kpi_rows:
        row_number += 1

        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        )

        value_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=float(value),
        )

        value_cell.font = FORMULA_FONT

        if label == "预计毛利率":
            value_cell.number_format = (
                PERCENT_FORMAT
            )
        else:
            value_cell.number_format = (
                MONEY_FORMAT
            )

    worksheet.cell(
        row=first_kpi_row,
        column=3,
        value=summary["account_count"],
    )

    worksheet.cell(
        row=first_kpi_row,
        column=4,
        value=summary[
            "transaction_count"
        ],
    )

    apply_table_borders(
        worksheet,
        start_row=first_kpi_row,
        end_row=row_number,
        column_count=4,
    )

    worksheet.freeze_panes = "A4"

    set_column_widths(
        worksheet,
        [
            28,
            22,
            16,
            16,
        ],
    )


def write_order_sheet(
    workbook,
    *,
    dashboard_data,
):
    worksheet = workbook.create_sheet(
        "订单明细"
    )

    worksheet.sheet_view.showGridLines = False

    headers = [
        "订单号",
        "医院",
        "工厂",
        "Invoice 编号",
        "PO 编号",
        "销售额 (EUR)",
        "采购额 (EUR)",
        "预计毛利润 (EUR)",
        "已收 (EUR)",
        "已付 (EUR)",
        "待收 (EUR)",
        "待付 (EUR)",
    ]

    set_sheet_title(
        worksheet,
        "按订单汇总的财务明细",
        end_column=len(headers),
    )

    header_row = 3

    for index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
        len(headers),
    )

    row_number = header_row

    for row in dashboard_data[
        "order_rows"
    ]:
        row_number += 1

        values = [
            row["order_number"],
            row["hospital_name"],
            row["factory_name"],
            row[
                "invoice_number_display"
            ],
            row["po_number_display"],
            float(row["sales"]),
            float(row["purchases"]),
            float(row["gross_profit"]),
            float(row["received"]),
            float(row["paid"]),
            float(
                row[
                    "receivable_remaining"
                ]
            ),
            float(
                row[
                    "payable_remaining"
                ]
            ),
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

        for column in range(
            6,
            13,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
            ).number_format = MONEY_FORMAT

    apply_table_borders(
        worksheet,
        start_row=header_row + 1,
        end_row=row_number,
        column_count=len(headers),
    )

    worksheet.freeze_panes = "A4"

    if row_number >= header_row:
        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"L{row_number}"
        )

    set_column_widths(
        worksheet,
        [
            22,
            28,
            28,
            24,
            24,
            16,
            16,
            19,
            16,
            16,
            16,
            16,
        ],
    )


def write_accounts_sheet(
    workbook,
    *,
    accounts,
    posted_by_account,
):
    worksheet = workbook.create_sheet(
        "应收应付"
    )

    worksheet.sheet_view.showGridLines = False

    headers = [
        "订单号",
        "正式文档编号",
        "方向",
        "医院 / 工厂",
        "开立日期",
        "截止日期",
        "币种",
        "原始金额",
        "有效收付金额",
        "剩余金额",
        "结算状态",
    ]

    set_sheet_title(
        worksheet,
        "应收与应付账户明细",
        end_column=len(headers),
    )

    header_row = 3

    for index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
        len(headers),
    )

    row_number = header_row

    for account in accounts:
        document = account.document
        order = document.order

        posted_amount = money(
            posted_by_account[
                account.id
            ]
        )

        remaining_amount = money(
            account.original_amount
            - posted_amount
        )

        if remaining_amount < ZERO:
            remaining_amount = ZERO

        row_number += 1

        values = [
            str(
                order.bon_de_commande
            ),
            document.document_number,
            account.get_direction_display(),
            account.counterparty_name,
            account.issue_date,
            account.due_date,
            account.currency,
            float(
                account.original_amount
            ),
            float(posted_amount),
            float(remaining_amount),
            account.get_status_display(),
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

        worksheet.cell(
            row=row_number,
            column=5,
        ).number_format = DATE_FORMAT

        worksheet.cell(
            row=row_number,
            column=6,
        ).number_format = DATE_FORMAT

        for column in range(
            8,
            11,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
            ).number_format = MONEY_FORMAT

    apply_table_borders(
        worksheet,
        start_row=header_row + 1,
        end_row=row_number,
        column_count=len(headers),
    )

    worksheet.freeze_panes = "A4"

    if row_number >= header_row:
        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"K{row_number}"
        )

    set_column_widths(
        worksheet,
        [
            22,
            25,
            16,
            30,
            14,
            14,
            10,
            17,
            18,
            17,
            18,
        ],
    )


def write_transactions_sheet(
    workbook,
    *,
    transactions,
):
    worksheet = workbook.create_sheet(
        "收付款流水"
    )

    worksheet.sheet_view.showGridLines = False

    headers = [
        "收付日期",
        "订单号",
        "正式文档编号",
        "方向",
        "医院 / 工厂",
        "金额",
        "币种",
        "方式",
        "参考号",
        "流水状态",
        "创建人",
        "创建时间",
        "冲销时间",
        "冲销原因",
        "备注",
    ]

    set_sheet_title(
        worksheet,
        "收款、付款与冲销流水",
        end_column=len(headers),
    )

    header_row = 3

    for index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
        len(headers),
    )

    row_number = header_row

    for transaction in transactions:
        account = transaction.account
        document = account.document
        order = document.order

        row_number += 1

        values = [
            transaction.payment_date,
            str(
                order.bon_de_commande
            ),
            document.document_number,
            account.get_direction_display(),
            account.counterparty_name,
            float(transaction.amount),
            account.currency,
            transaction.get_method_display(),
            transaction.reference,
            transaction.get_status_display(),
            user_display(
                transaction.created_by
            ),
            excel_datetime(
                transaction.created_at
            ),
            excel_datetime(
                transaction.reversed_at
            ),
            transaction.reversal_reason,
            transaction.notes,
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

        worksheet.cell(
            row=row_number,
            column=1,
        ).number_format = DATE_FORMAT

        worksheet.cell(
            row=row_number,
            column=6,
        ).number_format = MONEY_FORMAT

        worksheet.cell(
            row=row_number,
            column=12,
        ).number_format = DATETIME_FORMAT

        worksheet.cell(
            row=row_number,
            column=13,
        ).number_format = DATETIME_FORMAT

    apply_table_borders(
        worksheet,
        start_row=header_row + 1,
        end_row=row_number,
        column_count=len(headers),
    )

    worksheet.freeze_panes = "A4"

    if row_number >= header_row:
        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"O{row_number}"
        )

    set_column_widths(
        worksheet,
        [
            14,
            22,
            25,
            16,
            30,
            16,
            10,
            18,
            22,
            16,
            20,
            20,
            20,
            32,
            38,
        ],
    )


def build_finance_export_xlsx(
    *,
    reporting_currency="EUR",
    date_from=None,
    date_to=None,
    hospital_query="",
    factory_query="",
    order_query="",
    status="",
):
    filters = {
        "date_from": date_from,
        "date_to": date_to,
        "hospital_query": (
            hospital_query
        ),
        "factory_query": (
            factory_query
        ),
        "order_query": order_query,
        "status": status,
    }

    dashboard_data = (
        build_settlement_finance_dashboard_data(
            reporting_currency=(
                reporting_currency
            ),
            **filters,
        )
    )

    accounts = list(
        build_filtered_account_queryset(
            reporting_currency=(
                reporting_currency
            ),
            **filters,
        )
    )

    account_ids = [
        account.id
        for account in accounts
    ]

    transactions = list(
        PaymentTransaction.objects
        .filter(
            account_id__in=account_ids
        )
        .select_related(
            "account",
            "account__document",
            "account__document__order",
            "created_by",
        )
        .order_by(
            "payment_date",
            "id",
        )
    )

    posted_by_account = defaultdict(
        lambda: ZERO
    )

    for transaction in transactions:
        if (
            transaction.status
            == PaymentTransaction
            .Status
            .POSTED
        ):
            posted_by_account[
                transaction.account_id
            ] += money(
                transaction.amount
            )

    workbook = Workbook()

    workbook.properties.creator = (
        "Acoeur"
    )

    workbook.properties.title = (
        "Acoeur Finance Export"
    )

    workbook.properties.subject = (
        "Settlement finance data"
    )

    write_summary_sheet(
        workbook,
        dashboard_data=dashboard_data,
        filters=filters,
    )

    write_order_sheet(
        workbook,
        dashboard_data=dashboard_data,
    )

    write_accounts_sheet(
        workbook,
        accounts=accounts,
        posted_by_account=(
            posted_by_account
        ),
    )

    write_transactions_sheet(
        workbook,
        transactions=transactions,
    )

    output = BytesIO()
    workbook.save(output)

    timestamp = (
        timezone.localtime()
        .strftime("%Y%m%d_%H%M%S")
    )

    filename = (
        "acoeur_finance_"
        f"{timestamp}.xlsx"
    )

    return {
        "content": output.getvalue(),
        "filename": filename,
        "account_count": len(accounts),
        "transaction_count": len(
            transactions
        ),
    }
